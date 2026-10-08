"""Offline regressions: temporary storage only, no model or speech requests."""
import asyncio
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

API_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(API_DIR))
_runtime = tempfile.TemporaryDirectory()
os.environ.update({
    "DATABASE_URL": f"sqlite:///{Path(_runtime.name).as_posix()}/app.db",
    "UPLOAD_DIR": str(Path(_runtime.name) / "uploads"),
    "AUDIO_DIR": str(Path(_runtime.name) / "audio"),
    "LLM_API_KEY": "",
})

import fitz
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

import db as database
from config import settings
from db import Base, get_db
from main import app
from models.chapter import Chapter
from models.textbook import Textbook
from services import generator, tts


class APIReliabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.engine = create_engine(
            f"sqlite:///{(self.root / 'test.db').as_posix()}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.uploads = self.root / "uploads"
        self.audio = self.root / "audio"
        self.audio.mkdir()
        self.patches = [
            patch.object(settings, "upload_dir", str(self.uploads)),
            patch.object(settings, "audio_dir", str(self.audio)),
            patch.object(settings, "max_upload_size_mb", 1),
        ]
        for item in self.patches:
            item.start()

        def session_override():
            with self.sessions() as session:
                yield session

        app.dependency_overrides[get_db] = session_override
        self.client = TestClient(app)
        with self.sessions() as session:
            book = Textbook(title="测试教材", filename="source.pdf")
            session.add(book)
            session.flush()
            chapter = Chapter(
                textbook_id=book.id, title="测试章节", order=1, content="原文内容",
                generated_content="旧讲解", audio_filename="old.mp3",
                quiz_data='[{"question":"旧题"}]',
                knowledge_graph_data='{"nodes":[],"edges":[]}',
            )
            session.add(chapter)
            session.commit()
            self.book_id, self.chapter_id = book.id, chapter.id
        self.path = f"/api/textbooks/{self.book_id}/chapters/{self.chapter_id}"
        (self.audio / "old.mp3").write_bytes(b"old-audio")
        (self.audio / "old.json").write_text('[{"start":0,"end":1,"text":"旧字幕"}]', encoding="utf-8")

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        for item in reversed(self.patches):
            item.stop()
        self.engine.dispose()
        self.temp.cleanup()

    def chapter(self):
        return self.client.get(self.path).json()

    def advance_revision(self):
        with self.sessions() as session:
            chapter = session.get(Chapter, self.chapter_id)
            chapter.generated_content = "另一个请求的新讲解"
            chapter.generation_revision += 1
            session.commit()

    def test_regeneration_invalidates_all_artifacts_and_files(self):
        with patch("routers.generation.generate_chapter_content", return_value="新讲解"):
            response = self.client.post(self.path + "/generate", json={"style": "story"})
        self.assertEqual(response.status_code, 200)
        chapter = self.chapter()
        self.assertEqual(chapter["generated_content"], "新讲解")
        for key in ("audio_filename", "quiz_data", "knowledge_graph_data"):
            self.assertIsNone(chapter[key])
        self.assertEqual(list(self.audio.iterdir()), [])
        with self.sessions() as session:
            self.assertEqual(session.get(Chapter, self.chapter_id).generation_revision, 1)

    def test_failed_generation_preserves_existing_work(self):
        with patch("routers.generation.generate_chapter_content", side_effect=RuntimeError("offline")):
            response = self.client.post(self.path + "/generate", json={})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.chapter()["generated_content"], "旧讲解")
        self.assertEqual(self.chapter()["audio_filename"], "old.mp3")
        self.assertTrue((self.audio / "old.mp3").exists())

    def test_invalid_style_or_difficulty_rejected(self):
        for body in ({"style": "unknown"}, {"difficulty": "unknown"}):
            self.assertEqual(self.client.post(self.path + "/generate", json=body).status_code, 422)

    def test_conflicting_generation_cannot_overwrite_newer_content(self):
        def generate(**kwargs):
            self.advance_revision()
            return "过期讲解"
        with patch("routers.generation.generate_chapter_content", side_effect=generate):
            response = self.client.post(self.path + "/generate", json={})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.chapter()["generated_content"], "另一个请求的新讲解")

    def test_stale_quiz_and_graph_are_rejected(self):
        for endpoint, target, value in (
            ("quiz", "routers.quiz.generate_quiz", []),
            ("knowledge-graph", "routers.knowledge_graph.generate_knowledge_graph", {"nodes": [], "edges": []}),
        ):
            def generate(**kwargs):
                self.advance_revision()
                return value
            with patch(target, side_effect=generate):
                response = self.client.post(self.path + "/" + endpoint, json={})
            self.assertEqual(response.status_code, 409)
        self.assertEqual(self.chapter()["quiz_data"], '[{"question":"旧题"}]')

    def test_subtitles_and_audio_range_can_be_restored(self):
        response = self.client.get("/api/subtitles/old.json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()[0]["text"], "旧字幕")
        response = self.client.get("/api/audio/old.mp3", headers={"Range": "bytes=0-2"})
        self.assertEqual(response.status_code, 206)
        self.assertEqual(response.content, b"old")

    def test_audio_replacement_removes_old_files(self):
        async def fake_audio(content, output):
            Path(output).write_bytes(b"new-audio")
            Path(output).with_suffix(".json").write_text("[]")
            return output, [{"start": 0, "end": 1, "text": "新字幕"}]
        with patch("routers.tts.generate_audio", side_effect=fake_audio):
            response = self.client.post(self.path + "/tts")
        self.assertEqual(response.status_code, 200)
        self.assertFalse((self.audio / "old.mp3").exists())
        self.assertFalse((self.audio / "old.json").exists())
        self.assertEqual(len(list(self.audio.iterdir())), 2)

    def test_stale_audio_is_rejected_and_cleaned(self):
        async def fake_audio(content, output):
            Path(output).write_bytes(b"stale")
            Path(output).with_suffix(".json").write_text("[]")
            self.advance_revision()
            return output, []
        with patch("routers.tts.generate_audio", side_effect=fake_audio):
            response = self.client.post(self.path + "/tts")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(sorted(p.name for p in self.audio.iterdir()), ["old.json", "old.mp3"])

    def test_tts_service_failure_preserves_old_audio(self):
        with patch("routers.tts.generate_audio", side_effect=RuntimeError("offline")):
            response = self.client.post(self.path + "/tts")
        self.assertEqual(response.status_code, 502)
        self.assertEqual(self.chapter()["audio_filename"], "old.mp3")

    def test_upload_rejects_oversized_empty_and_malformed_files(self):
        for content, expected in ((b"x" * (1024 * 1024 + 1), 413), (b"", 400), (b"invalid", 400)):
            response = self.client.post("/api/upload", files={"file": ("book.pdf", content, "application/pdf")})
            self.assertEqual(response.status_code, expected)
            self.assertEqual(list(self.uploads.iterdir()), [])
        with self.sessions() as session:
            self.assertEqual(session.query(Textbook).count(), 1)

    def test_valid_pdf_at_size_limit_and_docx_upload(self):
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((72, 72), "Chapter 1\nSample textbook content")
            pdf = doc.tobytes()
        pdf += b"\n" * (1024 * 1024 - len(pdf))
        response = self.client.post("/api/upload", files={"file": ("Book.PDF", pdf, "application/pdf")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["title"], "Book")
        document = Document()
        document.add_heading("第一章", 1)
        document.add_paragraph("章节正文")
        buffer = io.BytesIO()
        document.save(buffer)
        response = self.client.post("/api/upload", files={"file": ("教材.docx", buffer.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["chapter_count"], 1)


class ServiceReliabilityTests(unittest.TestCase):
    def test_each_style_reaches_model_with_a_distinct_instruction(self):
        responses = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="讲解"))])
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kwargs: responses)))
        prompts = []
        def create(**kwargs):
            prompts.append(kwargs["messages"])
            return responses
        client.chat.completions.create = create
        with patch.object(generator, "get_client", return_value=client):
            for style in ("teacher", "concise", "story"):
                generator.generate_chapter_content("标题", "原文", "easy", style)
        self.assertEqual(len({messages[1]["content"] for messages in prompts}), 3)
        for messages in prompts:
            self.assertIn("原文", messages[1]["content"])
            self.assertIn("初学者", messages[1]["content"])

    def test_tts_uses_service_offsets_without_adding_gaps(self):
        calls = []
        class FakeCommunicate:
            def __init__(self, content, voice, **kwargs):
                calls.append((content, kwargs))
            async def stream(self):
                for offset in (0, 13_000_000, 26_000_000):
                    yield {"type": "audio", "data": b"audio"}
                    yield {"type": "SentenceBoundary", "offset": offset, "duration": 10_000_000, "text": "字幕"}
        with tempfile.TemporaryDirectory() as folder:
            output = str(Path(folder) / "test.mp3")
            with patch.object(tts.edge_tts, "Communicate", FakeCommunicate):
                _, subtitles = asyncio.run(tts.generate_audio("长文本" * 3000, output))
            self.assertEqual(len(calls), 1)
            self.assertEqual(calls[0][1]["boundary"], "SentenceBoundary")
            self.assertEqual([s["start"] for s in subtitles], [0, 1.3, 2.6])
            self.assertEqual(Path(output).read_bytes(), b"audio" * 3)
            self.assertEqual(json.loads(Path(output).with_suffix(".json").read_text(encoding="utf-8")), subtitles)

    def test_failed_stream_leaves_no_partial_audio(self):
        class BrokenCommunicate:
            def __init__(self, *args, **kwargs):
                pass
            async def stream(self):
                yield {"type": "audio", "data": b"partial"}
                raise RuntimeError("interrupted")
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(tts.edge_tts, "Communicate", BrokenCommunicate):
                with self.assertRaises(RuntimeError):
                    asyncio.run(tts.generate_audio("text", str(Path(folder) / "test.mp3")))
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_legacy_database_migration_preserves_content(self):
        with tempfile.TemporaryDirectory() as folder:
            engine = create_engine(f"sqlite:///{Path(folder).as_posix()}/legacy.db")
            with engine.begin() as connection:
                connection.execute(text('CREATE TABLE chapters (id INTEGER PRIMARY KEY, textbook_id INTEGER, title TEXT, "order" INTEGER, content TEXT, generated_content TEXT, audio_filename TEXT, quiz_data TEXT, knowledge_graph_data TEXT)'))
                connection.execute(text("INSERT INTO chapters (id, content) VALUES (1, 'preserved')"))
            with patch.object(database, "engine", engine):
                database.init_db()
                database.init_db()
            with engine.connect() as connection:
                self.assertEqual(connection.execute(text("SELECT content, generation_revision FROM chapters")).one(), ("preserved", 0))
            engine.dispose()


if __name__ == "__main__":
    unittest.main()
